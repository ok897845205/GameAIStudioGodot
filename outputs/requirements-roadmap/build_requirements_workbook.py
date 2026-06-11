from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


SOURCE = Path("outputs/requirements-roadmap/source.xlsx")
OUTPUT = Path("outputs/requirements-roadmap/0604_万象游戏创作平台roadmap_v1.2_需求拆解完成版.xlsx")

GENERATED_SHEETS = [
    "游戏创作平台roadmap细化版",
    "后续需求任务池",
    "版本里程碑",
    "风险与依赖",
    "字段说明",
]

REQ_ROWS = [
    {
        "id": "REQ-001",
        "module": "Godot 引擎与项目工程",
        "name": "Godot 引擎接入基础能力",
        "version": "万象 v0.1.0-alpha / 2026-06-10",
        "priority": "高",
        "status": "待验收",
        "scope": "Godot 4 项目识别/创建；2D/3D 模板复制与清理；运行时健康检查；最小 smoke 校验；阻塞原因展示。",
        "acceptance": "干净机器可创建 2D/3D 项目；Godot GUI 可打开；模板 smoke 返回明确结果；失败时 UI 展示原因和下一步。",
        "risk": "依赖内置 Godot 运行时、模板 export_presets、Windows 路径权限和打包资源目录。",
        "next": "补内测验收清单，跑 2D/3D 模板、缺引擎、缺导出预设、路径含中文/空格等用例。",
    },
    {
        "id": "REQ-002",
        "module": "AI CLI 与 Agent 运行",
        "name": "AI Coding 引擎扩展：Codex + Claude Code",
        "version": "万象 v0.1.0-alpha / 2026-06-10",
        "priority": "高",
        "status": "待验收",
        "scope": "保留现有对话入口；统一 ProviderAdapter；CLI 发现/健康检查/命令构建/流式输出/取消；文件变更回传。",
        "acceptance": "Codex、Claude Code 至少各完成一轮真实项目修改；账号未登录、PATH 缺失、额度不足能被识别并给出修复建议。",
        "risk": "本地 CLI 版本、账号登录态、KSCC 账号打通、Windows shim 和输出格式变化会影响稳定性。",
        "next": "建立 CLI 兼容矩阵，优先解决冲突问题和账号打通方案。",
    },
    {
        "id": "REQ-003",
        "module": "用户反馈闭环",
        "name": "用户使用反馈入口",
        "version": "万象 v0.1.0-alpha / 内测期",
        "priority": "中",
        "status": "规划中",
        "scope": "产品内反馈入口；一键联系/加群；自动携带版本、项目、Run、日志尾部、资源健康状态等上下文。",
        "acceptance": "内测用户能在 30 秒内提交问题；反馈记录可按问题/建议/失败/账号/导出分类；上下文足够研发复现。",
        "risk": "需要明确存储方式、隐私边界、是否上传日志、客服/社群承接方式。",
        "next": "先定反馈字段和回收渠道，再做轻入口；优先级低于 v0.2 核心链路。",
    },
    {
        "id": "REQ-004",
        "module": "账号体系与跨线协同",
        "name": "账号体系技术方案论证",
        "version": "研究 / 2026-06-20",
        "priority": "中",
        "status": "研究中",
        "scope": "创作工具与大赛网站账号关系；登录态、token、设备绑定、游客态、KSCC/外部 CLI 账号边界。",
        "acceptance": "输出技术方案文档，包含推荐路径、接口草案、风险、实施排期和 2026-07-10 大赛接入前置条件。",
        "risk": "大赛网站、账号中心、CLI 服务账号可能属于不同系统，协议和安全口径需要提前对齐。",
        "next": "约齐官网/账号/客户端负责人，对齐登录态传递与投稿 API。",
    },
    {
        "id": "REQ-005",
        "module": "素材与生成工具",
        "name": "文生图接入马良：手动导入优先",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "中",
        "status": "进行中",
        "scope": "保留马良画卷外部工具入口；支持用户生成图片后导入项目 assets；进入统一资源清单和素材槽位。",
        "acceptance": "用户能打开马良、拿到图片、导入当前 Godot 项目、在对话中被 Agent 识别并应用到场景或 UI。",
        "risk": "马良登录态、图片下载、跨域/文件保存、素材命名、重复导入和版权标记都要有明确规则。",
        "next": "v0.2 只做手动导入和绑定，自动生成/替换流程放到 9 月版本。",
    },
    {
        "id": "REQ-006",
        "module": "用户反馈闭环",
        "name": "用户反馈机制完善",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "低",
        "status": "规划中",
        "scope": "基于 6/10 内测问题建立标签、优先级、处理状态、FAQ 和痛点聚类。",
        "acceptance": "每条反馈可关联版本、项目、运行记录和处理状态；周度能输出 Top 问题和改进建议。",
        "risk": "如果没有统一数据结构，反馈会散落在群聊和手工表格里，难以转成研发任务。",
        "next": "随 v0.2 核心链路一起埋点，但功能入口可后置。",
    },
    {
        "id": "REQ-007",
        "module": "Godot 引擎与项目工程",
        "name": "Godot 引擎接入完善：真实工程闭环",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "高",
        "status": "进行中",
        "scope": "项目创建/打开/模板选择/健康检查；Agent 修改后触发 import、validate、Web 导出、预览和失败指引。",
        "acceptance": "用户从一句话到 Web 可玩 zip 的链路稳定；导出失败不会交付旧模板；错误能进入修复回合。",
        "risk": "Godot 导入慢、模板质量、Web 导出依赖、Agent 变更为空、项目路径权限都会影响闭环。",
        "next": "用真实游戏需求建立回归样例，覆盖 2D、3D、中文项目名、失败修复轮。",
    },
    {
        "id": "REQ-008",
        "module": "素材与生成工具",
        "name": "2D 资源管理",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "高",
        "status": "待方案",
        "scope": "资源搜索/导入/预览/分类/替换；AssetManifest；素材槽位；与 Agent 提示词和 Godot assets 目录打通。",
        "acceptance": "用户能把图片资源放入项目并绑定用途；Agent 能看到素材清单；替换素材不破坏场景引用。",
        "risk": "网页式搜索方案、版权来源、资源命名、重复文件、缩略图和项目引用路径需要定规则。",
        "next": "先出 2D 资源信息架构和最小导入方案，linli 负责项需确认任务边界。",
    },
    {
        "id": "REQ-009",
        "module": "素材与生成工具",
        "name": "3D 模型生成沿用混元",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "高",
        "status": "规划中",
        "scope": "保留混元 3D 生成机制；支持模型导入、格式校验、贴图归档、预览和替换；复杂自动化后置。",
        "acceptance": "至少 3 个 GLB/GLTF 模型能导入 3D 模板，材质不丢失，Web 导出后可加载。",
        "risk": "模型格式、贴图路径、体积、动画、Web 性能和许可证风险需要 QA 规则。",
        "next": "定义 3D 资产验收标准，先做导入/预览/校验，不重写生成链路。",
    },
    {
        "id": "REQ-010",
        "module": "核心创作链路",
        "name": "核心创作链路稳定化",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "高",
        "status": "进行中",
        "scope": "从用户想法到制作人/策划/程序/美术/QA，再到 Godot 校验、修复、Git 保存、Web 导出和预览。",
        "acceptance": "5 条真实游戏需求中至少 4 条能生成可玩的 Web 版本；失败案例有明确定位和可继续修复入口。",
        "risk": "Agent 不改文件、生成不可运行代码、QA 反馈不可执行、导出旧模板、上下文过长。",
        "next": "建立端到端回归任务，按失败类型补保护和修复回合。",
    },
    {
        "id": "REQ-011",
        "module": "运行可视化与诊断",
        "name": "AI 运行结果可视化",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "高",
        "status": "已具备基础能力",
        "scope": "右侧运行面板；每个 Agent/导出/校验/打包/预览步骤状态；流式输出；失败原因；日志和文件变更摘要。",
        "acceptance": "用户能看懂当前卡在哪一步；失败信息能定位到 CLI、Godot、导出、zip 或预览；支持取消运行。",
        "risk": "错误信息过长、日志噪声、多个项目并发、取消状态与后台进程一致性。",
        "next": "补失败分类文案和关键日志链接，减少“运行失败但不知道为什么”。",
    },
    {
        "id": "REQ-012",
        "module": "大赛投稿与跨线接口",
        "name": "工具到大赛网站投稿对接",
        "version": "万象 v0.2.0 / 对齐 2026-07-10 大赛",
        "priority": "中",
        "status": "依赖外部进度",
        "scope": "工具内生成作品后，对接官网投稿协议；上传 zip/元数据/截图/作者信息；展示提交状态和失败原因。",
        "acceptance": "用户在工具内完成作品后可跳转或直接投稿；官网能识别作品包和账号；失败可重试。",
        "risk": "官网接口、账号态、包体大小、审核字段、作品截图和隐私协议未确定。",
        "next": "先了解大赛官网投稿链路进度，锁定协议和最小字段。",
    },
    {
        "id": "REQ-013",
        "module": "账号体系与跨线协同",
        "name": "用户登录态/账号体系实施",
        "version": "万象 v0.2.0 / 方案后实施",
        "priority": "中",
        "status": "依赖方案",
        "scope": "根据 REQ-004 方案实现登录、会话保存、账号展示、投稿鉴权和退出登录。",
        "acceptance": "用户账号态能在工具内被识别；投稿或大赛相关能力无需重复登录；token 安全存储。",
        "risk": "安全合规、账号中心接口、自动更新后的配置迁移、离线使用体验。",
        "next": "等 2026-06-20 技术方案定稿后拆实施 issue。",
    },
    {
        "id": "REQ-014",
        "module": "研究里程碑",
        "name": "微信小游戏研究里程碑",
        "version": "研究 M1 / 2026-06-30",
        "priority": "中",
        "status": "研究中",
        "scope": "调研 Godot/网页游戏到微信小游戏路径；登录奖励、广告植入、审核限制、性能和包体限制。",
        "acceptance": "产出调研分析报告、规划讨论稿和技术可行性结论，明确是否进入后续版本实施。",
        "risk": "与主线 Godot Web 导出、Cocos 扩展和大赛节奏竞争资源。",
        "next": "先作为研究里程碑，不绑定 v0.2 发布。",
    },
    {
        "id": "REQ-015",
        "module": "多引擎扩展",
        "name": "Cocos 引擎接入",
        "version": "万象 v0.4.0 待定 / 约 2026-09-30",
        "priority": "高",
        "status": "规划占位",
        "scope": "在 Godot 闭环稳定后扩展 Cocos：模板、运行、预览、导出、Agent 上下文和资源目录规则。",
        "acceptance": "完成 Cocos POC：创建示例、运行预览、导出 Web 包、Agent 修改一次并通过校验。",
        "risk": "多引擎抽象过早会拖慢 v0.2；需要等 Godot Adapter 边界稳定后再抽。",
        "next": "先沉淀 EngineAdapter 接口草案，9 月前不抢核心链路资源。",
    },
    {
        "id": "REQ-016",
        "module": "移动端能力",
        "name": "移动端适配：横屏/竖屏",
        "version": "万象 v0.4.0 待定 / 约 2026-07 初步验证",
        "priority": "中",
        "status": "规划中",
        "scope": "移动 Web 预览、横竖屏配置、触控输入、安全区、分辨率适配和模板约束。",
        "acceptance": "至少一个 2D 模板在手机浏览器横屏/竖屏可操作；导出包说明包含移动适配要求。",
        "risk": "Godot Web 移动性能、输入映射、软键盘、安全区和微信小游戏路线耦合。",
        "next": "先做横竖屏技术验证，手机游戏样例后置。",
    },
    {
        "id": "REQ-017",
        "module": "样例与模板",
        "name": "手机端游戏样例",
        "version": "万象 v0.4.0 待定 / 约 2026-09-30",
        "priority": "中",
        "status": "规划占位",
        "scope": "补充适合手机端的可玩样例：轻操作、短局、触控友好、可用于大赛展示。",
        "acceptance": "至少 3 个手机端样例可创建、预览、导出，并作为 Agent 生成的参考模板。",
        "risk": "样例质量直接影响用户首感，需要玩法、美术和性能一起验收。",
        "next": "等移动端适配验证后确认样例类型。",
    },
    {
        "id": "REQ-018",
        "module": "素材与生成工具",
        "name": "Chaos config UI 打通",
        "version": "万象 v0.4.0 待定 / 约 2026-09-30",
        "priority": "中",
        "status": "规划占位",
        "scope": "从外部接入升级为产品内 config UI；支持参数配置、生成记录、资源回写和错误状态。",
        "acceptance": "用户无需离开创作工具即可配置并触发 Chaos 相关能力，结果进入资源管理体系。",
        "risk": "涉及外部服务 API、账号、成本、失败重试和 UI 复杂度，不能挤占 v0.2 马良手动导入。",
        "next": "先用马良/混元手动导入沉淀资源管线，再做完整 config UI。",
    },
    {
        "id": "REQ-019",
        "module": "样例与模板",
        "name": "热门小游戏样例库",
        "version": "规划占位 / 7-9 月滚动",
        "priority": "中",
        "status": "待调研",
        "scope": "参考榜单热门类型，沉淀可复用模板和 Agent 提示词样例，例如多人社交、派对、轻竞技、解谜等方向。",
        "acceptance": "每类样例有可玩 demo、玩法说明、资产清单、生成提示词和可扩展任务。",
        "risk": "需要规避直接复制热门游戏，同时控制样例研发成本。",
        "next": "先做市场/榜单拆解，再挑 2-3 类投入样例。",
    },
    {
        "id": "REQ-020",
        "module": "研发流程与质量",
        "name": "程序开发流程改造试点",
        "version": "内部改造 / 7-9 月滚动",
        "priority": "中",
        "status": "规划中",
        "scope": "把需求拆解、issue、TDD/回归、发布检查和 AI Agent 协作纳入固定流程。",
        "acceptance": "关键需求都有任务拆解、验收标准、测试命令和发布检查；失败可追踪到 issue。",
        "risk": "流程太重会影响迭代速度；需要先在 v0.2 核心链路试点。",
        "next": "从本表任务池开始，建立周度更新和完成定义。",
    },
    {
        "id": "REQ-021",
        "module": "发布与内测交付",
        "name": "内测发布、安装包与更新链路稳定化",
        "version": "万象 v0.2.0 / 2026-06-30",
        "priority": "高",
        "status": "已具备基础能力",
        "scope": "Windows 安装包、更新清单、可选/强制更新、更新日志、安装包校验、发布脚本和冒烟检查。",
        "acceptance": "内测包可安装、可检查更新、可升级到新版本；发布前 verify:release 通过。",
        "risk": "服务器配置、HTTPS、安装包签名/缓存、latest.yml 与 update.json 一致性。",
        "next": "把发布流程写成验收清单，v0.2 每次发包前执行。",
    },
]

TASK_ROWS = [
    ["T001", "REQ-001", "v0.1-alpha", "验收", "高", "Godot 运行时体检清单", "列出引擎目录、GUI/控制台可执行文件、版本探测、模板目录、export_presets 必检项。", "体检清单 + UI 状态口径", "缺任一项时有明确错误、修复建议和日志路径。", "内置 engine、resourceRoot", "客户端/QA", "待验收", "2026-06-12", ""],
    ["T002", "REQ-001", "v0.1-alpha", "验收", "高", "2D/3D 模板复制清理验证", "验证复制新项目时排除 .godot、build、dist 等历史产物。", "模板复制验收记录", "新项目只包含干净模板源码，项目能被 Godot 打开。", "模板目录", "客户端/QA", "待验收", "2026-06-12", ""],
    ["T003", "REQ-001", "v0.1-alpha", "测试", "高", "模板 smoke 自动化用例", "执行 2D/3D validate_project 与 Web 导出预设检查。", "smoke:templates 结果", "成功/失败都有可读输出；失败不进入后续导出。", "Godot console", "客户端/QA", "待验收", "2026-06-13", ""],
    ["T004", "REQ-001", "v0.2", "增强", "中", "路径兼容场景覆盖", "覆盖中文路径、空格路径、无权限目录、安装目录误选等情况。", "路径兼容测试表", "所有路径错误能阻断并提示，不破坏用户项目。", "Windows 文件系统", "客户端/QA", "待开始", "2026-06-18", ""],
    ["T005", "REQ-002", "v0.1-alpha", "验收", "高", "CLI 兼容矩阵", "整理 Codex、Claude Code、KSCC、Kimi 的命令、版本、登录态、headless probe、环境变量。", "CLI 兼容矩阵", "每个 CLI 至少有发现、健康检查、未登录提示和安装建议。", "本地 CLI 安装", "客户端/AI", "进行中", "2026-06-14", "重点解决冲突问题。"],
    ["T006", "REQ-002", "v0.1-alpha", "验收", "高", "Codex/Claude 真实回合验证", "用同一 Godot 项目分别跑 Codex、Claude Code，验证流式输出、文件变更、完成状态。", "双 CLI 验收记录", "每个 CLI 至少完成一次脚本或场景改动并被 UI 捕获。", "账号登录态", "客户端/QA", "待验收", "2026-06-14", ""],
    ["T007", "REQ-002", "v0.2", "增强", "高", "认证失败分类文案", "把 401、未登录、额度、命令不存在、权限拒绝等失败分类到统一提示。", "错误分类表 + UI 文案", "用户能按提示完成 login 或切换 CLI。", "CLI 输出格式", "客户端/产品", "待开始", "2026-06-20", ""],
    ["T008", "REQ-002", "v0.2", "增强", "中", "输出解析与完成判定收口", "统一判断 Agent 是否真正完成、是否产生文件变更、是否应进入导出。", "完成判定规则", "无文件变更时不交付旧模板，并提示继续修改。", "AgentService/WorkflowService", "客户端", "进行中", "2026-06-21", ""],
    ["T009", "REQ-003", "v0.1-alpha", "产品", "中", "反馈入口位置确认", "决定反馈入口放在顶部、关于弹层、运行失败提示还是项目面板。", "反馈入口交互稿", "内测用户能在不打断创作的情况下提交反馈。", "产品设计", "产品/设计", "待开始", "2026-06-17", "优先级低于核心链路。"],
    ["T010", "REQ-003", "v0.2", "方案", "中", "反馈上下文字段定义", "定义 appVersion、projectId、runId、agent、CLI、Godot 状态、日志尾部、用户描述等字段。", "反馈 payload 字段表", "字段能支持研发定位 80% 运行失败。", "日志隐私边界", "产品/客户端", "待开始", "2026-06-18", ""],
    ["T011", "REQ-003", "v0.2", "方案", "低", "反馈存储与回收渠道", "确定本地存储、服务器、表单、飞书/群渠道的 MVP 路径。", "反馈回收方案", "每条反馈能追踪处理状态。", "后端/运营资源", "产品/后端/运营", "待开始", "2026-06-20", ""],
    ["T012", "REQ-004", "研究", "研究", "中", "账号系统边界调研", "梳理创作工具账号、大赛网站账号、AI CLI 账号、KSCC 账号的边界和打通必要性。", "账号边界图", "明确哪些账号必须统一，哪些只做提示。", "官网/账号团队", "产品/架构", "研究中", "2026-06-15", ""],
    ["T013", "REQ-004", "研究", "方案", "中", "登录态传递方案", "比较 SSO、网页登录授权、token 注入、设备码、跳转投稿等方案。", "技术方案文档", "推荐路径含安全、开发量、用户体验和风险。", "账号中心接口", "架构/客户端/后端", "研究中", "2026-06-20", ""],
    ["T014", "REQ-004", "研究", "协同", "中", "大赛前置接口对齐会", "与大赛官网确认投稿、用户、作品包、审核字段和截止时间。", "会议纪要 + API 待办", "形成可实施接口列表。", "大赛官网进度", "产品/官网/客户端", "待开始", "2026-06-20", ""],
    ["T015", "REQ-005", "v0.2", "实现", "中", "马良入口可用性验证", "验证内嵌马良画卷 webview 的登录态、加载失败、网络失败和刷新行为。", "入口验收记录", "能稳定打开并保留登录态；失败可提示用户。", "马良服务", "客户端/QA", "进行中", "2026-06-18", ""],
    ["T016", "REQ-005", "v0.2", "实现", "中", "图片导入当前项目", "支持用户选择或拖入马良生成图片，复制到当前 Godot 项目 assets 目录。", "导入功能", "导入后文件名安全、路径可预览、Agent 上下文可见。", "项目文件预览/附件逻辑", "客户端", "待开始", "2026-06-23", ""],
    ["T017", "REQ-005", "v0.2", "设计", "中", "素材槽位最小模型", "定义角色、背景、UI、道具等 2D 素材用途字段，避免只是一堆文件。", "AssetSlot 字段草案", "导入资源能被标注用途并用于提示词。", "2D 资源管理", "产品/客户端", "待开始", "2026-06-24", ""],
    ["T018", "REQ-005", "v0.2", "测试", "中", "马良导入到 Godot 场景闭环", "导入一张图后让 Agent 应用到 HUD 或角色占位资源。", "端到端测试记录", "Web 预览中能看到导入素材。", "Agent 能力/Godot 导出", "QA/客户端", "待开始", "2026-06-27", ""],
    ["T019", "REQ-006", "v0.2", "运营", "低", "内测问题标签体系", "建立问题、建议、账号、导出、素材、性能、崩溃、体验等标签。", "反馈标签表", "每条反馈可落到一个主标签和一个优先级。", "反馈入口", "产品/运营", "待开始", "2026-06-25", ""],
    ["T020", "REQ-006", "v0.2", "运营", "低", "Top 问题周报模板", "把内测反馈聚类为周报：现象、影响范围、复现、负责人、状态。", "周报模板", "能转成研发任务池。", "反馈数据", "产品/运营", "待开始", "2026-06-28", ""],
    ["T021", "REQ-007", "v0.2", "实现", "高", "Agent 修改后 Godot import/validate", "Agent 改动脚本、场景、资源后自动或手动触发项目校验。", "校验链路", "失败时进入修复回合并保留输出。", "Godot console", "客户端", "进行中", "2026-06-20", ""],
    ["T022", "REQ-007", "v0.2", "实现", "高", "Web 导出产物检查", "检查 index.html、wasm、pck、manifest，并记录大小和路径。", "产物检查结果", "缺关键文件时阻断 zip 打包。", "ExportService", "客户端/QA", "已具备基础能力", "2026-06-21", "继续补真实样例。"],
    ["T023", "REQ-007", "v0.2", "实现", "高", "预览服务失败兜底", "Web 构建缺失、端口占用、旧缓存、路径越界时提示明确。", "预览错误处理", "不会展示陈旧构建；错误可复现。", "PreviewServer", "客户端", "进行中", "2026-06-22", ""],
    ["T024", "REQ-007", "v0.2", "测试", "高", "真实工程端到端回归", "选 5 条游戏需求，覆盖 2D/3D、素材、脚本、UI、导出。", "回归报告", "至少 4 条成功生成可玩 Web 版本。", "AI CLI/Godot", "QA/产品/客户端", "待开始", "2026-06-28", ""],
    ["T025", "REQ-008", "v0.2", "方案", "高", "2D 资源信息架构", "定义资源类型、来源、用途、标签、预览图、项目路径、引用关系。", "2D 资源模型方案", "能覆盖马良导入、用户本地导入、样例资源。", "AssetManifest", "产品/客户端（linli 待确认）", "待方案", "2026-06-17", ""],
    ["T026", "REQ-008", "v0.2", "方案", "高", "资源搜索/网页式方案评估", "评估网页搜索、内置资源库、外部资源站、用户本地导入的 MVP 可行性。", "方案对比表", "明确 v0.2 采用哪条路径。", "版权/网络/资源来源", "产品/设计/客户端", "待方案", "2026-06-19", ""],
    ["T027", "REQ-008", "v0.2", "实现", "高", "本地 2D 资源导入 MVP", "支持 PNG/JPG/WebP 导入到 assets/images，并生成基础元数据。", "导入 MVP", "导入资源能预览、重命名、删除或替换。", "文件预览服务", "客户端", "待开始", "2026-06-24", ""],
    ["T028", "REQ-008", "v0.2", "实现", "高", "Agent 上下文注入资源清单", "把导入资源和用途写入 Agent 可读上下文，减少 Agent 不知道素材位置。", "资源上下文段落", "Agent 回复能引用正确 res:// 路径。", "AgentContextService", "客户端/AI", "待开始", "2026-06-25", ""],
    ["T029", "REQ-008", "v0.2", "测试", "中", "资源替换不破坏引用", "验证替换同用途图片后 Godot 场景引用不丢。", "替换测试用例", "Web 预览显示替换后的资源。", "Godot import", "QA/客户端", "待开始", "2026-06-28", ""],
    ["T030", "REQ-009", "v0.2", "方案", "高", "混元 3D 现有机制梳理", "明确当前混元生成/下载/导入路径和可复用边界。", "3D 生成机制梳理", "v0.2 不重复造生成链路，只保留接入点。", "混元服务", "产品/算法/客户端", "待开始", "2026-06-18", ""],
    ["T031", "REQ-009", "v0.2", "实现", "高", "3D 模型导入与格式校验", "支持 GLB/GLTF 及贴图目录归档，校验文件体积和资源引用。", "模型导入 MVP", "至少 3 个模型导入后材质不丢。", "Godot 3D 模板", "客户端/QA", "待开始", "2026-06-25", ""],
    ["T032", "REQ-009", "v0.2", "测试", "中", "3D Web 导出性能检查", "记录模型大小、加载时间、Web 导出是否成功。", "3D 资源 QA 表", "超限资源有提示，不阻塞普通模型使用。", "Web 导出", "QA/客户端", "待开始", "2026-06-29", ""],
    ["T033", "REQ-010", "v0.2", "设计", "高", "端到端成功定义", "定义一句话生成游戏的成功标准：可运行、可玩、可预览、可导出、有版本记录。", "成功标准文档", "产品、研发、QA 对完成定义一致。", "核心链路", "产品/QA/客户端", "待开始", "2026-06-15", ""],
    ["T034", "REQ-010", "v0.2", "实现", "高", "质量修复轮策略", "QA 反馈或 Godot 校验失败时，把上下文交给程序 Agent 修复并复检。", "修复轮策略", "修复后再次校验，仍失败时展示原因。", "WorkflowService", "客户端/AI", "已具备基础能力", "2026-06-20", "继续做真实场景调优。"],
    ["T035", "REQ-010", "v0.2", "实现", "高", "无文件变更保护", "当 Agent 没有产生项目文件变更时，跳过导出/zip/预览。", "保护逻辑", "不会把旧模板误交付给用户。", "Git/File change", "客户端", "已具备基础能力", "2026-06-20", ""],
    ["T036", "REQ-010", "v0.2", "测试", "高", "五类游戏需求回归集", "准备平台跳跃、射击、解谜、跑酷、3D 收集等真实需求。", "回归 prompt 集", "每条 prompt 有期望玩法和验收点。", "产品/QA", "产品/QA", "待开始", "2026-06-22", ""],
    ["T037", "REQ-010", "v0.2", "测试", "高", "v0.2 发布前端到端演练", "完整跑创建、团队工作流、修复、Git、导出、预览、zip。", "发布演练报告", "阻断问题清零或有明确降级方案。", "全部核心能力", "QA/客户端", "待开始", "2026-06-29", ""],
    ["T038", "REQ-011", "v0.2", "优化", "高", "运行时间线信息密度优化", "让用户区分 Agent、校验、导出、打包、预览每一步状态。", "运行面板优化清单", "卡住时用户能知道正在等什么。", "RunService/UI", "设计/客户端", "进行中", "2026-06-21", ""],
    ["T039", "REQ-011", "v0.2", "优化", "中", "失败原因与下一步指引", "把 CLI/Godot/导出/预览/Git 错误转成用户可执行建议。", "错误文案库", "常见失败不只显示 stderr。", "错误分类", "产品/客户端", "待开始", "2026-06-24", ""],
    ["T040", "REQ-011", "v0.2", "实现", "中", "日志与证据入口", "在失败步骤提供 app.log、project.log、Agent 上下文、变更文件入口。", "证据入口", "研发能快速定位失败上下文。", "日志服务/文件预览", "客户端/QA", "待开始", "2026-06-26", ""],
    ["T041", "REQ-012", "v0.2", "调研", "中", "大赛投稿字段清单", "确认作品名、作者、账号、简介、标签、截图、zip、manifest、版本等字段。", "投稿字段表", "字段与官网表单/接口一致。", "大赛官网", "产品/官网", "待开始", "2026-06-18", ""],
    ["T042", "REQ-012", "v0.2", "方案", "中", "投稿方式决策", "选择直接 API 投稿、工具跳转官网带包、或导出后人工上传。", "投稿 MVP 方案", "2026-07-10 前可落地。", "账号方案/官网进度", "产品/架构", "待开始", "2026-06-24", ""],
    ["T043", "REQ-012", "v0.2", "实现", "中", "作品包 manifest 对齐", "确保 zip 内 gameaistudio-export.json 满足官网识别需求。", "manifest 字段调整", "官网能读取作品元数据。", "ExportService/官网", "客户端/官网", "待开始", "2026-06-28", ""],
    ["T044", "REQ-013", "v0.2", "实现", "中", "工具内账号态展示", "根据方案展示已登录/未登录、账号昵称、投稿权限和退出入口。", "账号状态 UI", "用户知道当前投稿账号。", "账号方案", "客户端/设计", "依赖方案", "2026-06-30", ""],
    ["T045", "REQ-013", "v0.2", "实现", "中", "token 安全存储", "确定并实现 token 存储、刷新、清除、迁移和日志脱敏。", "安全存储实现", "日志不会泄露 token，重启后可保持登录。", "账号中心", "客户端/安全", "依赖方案", "2026-07-03", ""],
    ["T046", "REQ-014", "研究", "调研", "中", "微信小游戏平台限制调研", "梳理包体、性能、审核、登录、广告、支付/奖励等限制。", "调研分析报告", "列出对当前 Godot Web 链路的影响。", "微信平台文档", "产品/技术", "研究中", "2026-06-24", ""],
    ["T047", "REQ-014", "研究", "技术", "中", "Godot 到微信小游戏可行性", "验证 Godot Web 产物是否可转小游戏，或是否需要 Cocos/其他路径。", "技术验证结论", "给出推荐方案和不可行点。", "Godot/Cocos", "架构/客户端", "研究中", "2026-06-28", ""],
    ["T048", "REQ-014", "研究", "规划", "低", "微信小游戏规划讨论稿", "明确是否进入 7-9 月版本，以及与移动端/账号/广告的关系。", "规划讨论稿", "能用于立项或放弃决策。", "研究结论", "产品/架构", "待开始", "2026-06-30", ""],
    ["T049", "REQ-015", "v0.4", "技术", "中", "EngineAdapter 接口草案", "从 Godot 抽象创建、打开、校验、导出、预览、资源路径等接口。", "接口草案", "不影响 v0.2 交付，可支撑 Cocos POC。", "Godot 链路稳定", "架构/客户端", "待开始", "2026-07-15", ""],
    ["T050", "REQ-015", "v0.4", "POC", "高", "Cocos 模板 POC", "准备最小 Cocos 项目模板，验证创建和预览。", "Cocos POC", "能创建、运行、导出 Web。", "Cocos 环境", "客户端", "规划占位", "2026-09-15", ""],
    ["T051", "REQ-016", "v0.4", "技术", "中", "横竖屏配置验证", "验证 Godot Web 横屏/竖屏、分辨率、缩放和安全区。", "移动适配验证", "手机浏览器可操作且画面不裁切关键 UI。", "Godot Web", "客户端/QA", "待开始", "2026-07-05", ""],
    ["T052", "REQ-016", "v0.4", "设计", "中", "触控输入规范", "定义虚拟摇杆、点击、滑动、按钮尺寸和 HUD 安全区。", "触控规范", "样例游戏按统一规则实现。", "移动适配验证", "设计/客户端", "待开始", "2026-07-12", ""],
    ["T053", "REQ-017", "v0.4", "内容", "中", "手机端样例选型", "选择 3 个轻量样例方向并写出玩法边界。", "样例选题表", "每个样例有一句话玩法和素材需求。", "市场/榜单调研", "产品/策划", "规划占位", "2026-08-01", ""],
    ["T054", "REQ-017", "v0.4", "内容", "中", "手机样例模板制作", "实现 3 个手机端模板并接入创建流程。", "样例模板", "可创建、预览、导出。", "移动适配", "客户端/策划/美术", "规划占位", "2026-09-25", ""],
    ["T055", "REQ-018", "v0.4", "方案", "中", "Chaos config UI 范围定义", "确认哪些参数产品内配置，哪些仍跳外部工具。", "config UI 方案", "MVP 不扩大到完整自动生成平台。", "Chaos 服务 API", "产品/设计/算法", "规划占位", "2026-08-15", ""],
    ["T056", "REQ-018", "v0.4", "实现", "中", "生成记录与资源回写", "保存 Chaos 生成记录，并把结果写回资源管理。", "记录/回写 MVP", "生成结果可追踪、可复用。", "资源管理", "客户端/算法", "规划占位", "2026-09-20", ""],
    ["T057", "REQ-019", "规划", "调研", "中", "热门小游戏类型拆解", "拆榜单游戏的核心循环、操作、关卡、美术要求和可模板化程度。", "类型调研表", "选出 2-3 个适合大赛和 AI 生成的方向。", "市场调研", "产品/策划", "待调研", "2026-07-20", ""],
    ["T058", "REQ-019", "规划", "内容", "中", "样例提示词库", "为每类样例沉淀制作人/策划/程序/美术/QA 提示词。", "提示词库", "Agent 能按样例稳定生成变体。", "核心链路稳定", "产品/AI/策划", "待开始", "2026-08-15", ""],
    ["T059", "REQ-020", "内部", "流程", "中", "任务池更新节奏", "把本表作为周会跟踪基线，维护状态、负责人、阻塞和下一步。", "周更机制", "每周至少更新一次状态和阻塞。", "团队执行", "产品/项目管理", "待开始", "2026-06-14", ""],
    ["T060", "REQ-020", "内部", "质量", "中", "发布前检查门禁", "把 typecheck、test、smoke:webzip、smoke:templates、dist:win 纳入发布清单。", "发布检查清单", "v0.2 发包前按清单执行并记录结果。", "CI/本地环境", "客户端/QA", "已具备基础能力", "2026-06-28", ""],
    ["T061", "REQ-021", "v0.2", "发布", "高", "安装包 smoke", "安装、启动、创建项目、运行一次团队工作流、检查更新入口。", "安装包验收记录", "旧版本升级和新装都可用。", "Windows installer", "QA/客户端", "待开始", "2026-06-29", ""],
    ["T062", "REQ-021", "v0.2", "发布", "高", "update.json/latest.yml 一致性检查", "检查版本号、sha512、sha256、包大小、releaseNotes 和固定下载链接。", "更新清单检查表", "关于弹层和 electron-updater 读取一致。", "发布服务器", "客户端/运维", "待开始", "2026-06-29", ""],
    ["T063", "REQ-021", "v0.2", "发布", "中", "强制更新策略回归", "验证 patch 可选、minor/major 强制、minSupportedVersion 阻断。", "更新策略回归记录", "强制更新时阻止创建/运行/发送 Agent。", "UpdateService", "QA/客户端", "待开始", "2026-06-29", ""],
]

MILESTONES = [
    ["M1", "2026-06-10", "万象 v0.1.0-alpha", "核心功能内测验证", "Godot 基础能力、AI CLI 扩展、反馈入口方案、账号方案启动。", "是否能完成真实 Godot 项目创建和一次 AI 修改。", "内测可跑；关键阻塞记录到任务池。"],
    ["M2", "2026-06-20", "账号/跨线研究", "账号体系技术方案定稿", "创作工具与大赛网站账号、投稿鉴权、KSCC/CLI 账号边界。", "是否支持 2026-07-10 大赛打通。", "方案文档、接口草案、实施任务拆出。"],
    ["M3", "2026-06-30", "万象 v0.2.0", "正式版核心闭环", "Godot 完整支持、核心创作链路稳定、马良手动导入、2D/3D 资源管理 MVP、反馈机制完善。", "是否从演示跑通升级到真能用。", "5 条真实需求回归，至少 4 条可玩且可导出。"],
    ["M4", "2026-06-30", "研究 M1", "微信小游戏研究里程碑", "微信小游戏调研报告、规划讨论稿、技术可行性结论。", "是否进入后续版本。", "报告评审通过，明确实施/暂缓。"],
    ["M5", "2026-07-10", "大赛接入节点", "工具到大赛投稿链路", "作品包、manifest、账号态、官网接口或跳转投稿。", "是否能服务大赛投稿。", "至少完成最小投稿路径或明确人工替代方案。"],
    ["M6", "2026-07 初", "移动端技术验证", "横屏/竖屏适配验证", "移动 Web 预览、触控输入、安全区和分辨率。", "是否支持手机端样例投入。", "1 个模板手机可操作，问题清单明确。"],
    ["M7", "2026-09-30", "万象 v0.4.0 待定", "多引擎与生成工具升级", "Cocos POC、Chaos config UI、手机端样例、热门小游戏样例库。", "Godot 主线是否已经稳定。", "POC 或规划项达到可立项标准。"],
]

RISKS = [
    ["R001", "AI CLI", "本地 CLI 登录态/账号冲突导致 Agent 无法执行。", "高", "中", "建立 CLI 兼容矩阵；headless probe；错误分类文案；账号方案对齐。", "客户端/账号", "2026-06-20"],
    ["R002", "Godot", "Godot Web 导出失败或交付旧模板，用户误以为生成成功。", "高", "中", "导出前必须检查文件变更；导出后检查 wasm/pck/html；失败阻断 zip。", "客户端/QA", "2026-06-21"],
    ["R003", "资源", "2D 资源管理方案不清，导致马良/本地导入/Agent 使用互相割裂。", "高", "高", "先定 AssetManifest 和 AssetSlot 最小模型，再做导入。", "产品/客户端", "2026-06-19"],
    ["R004", "3D 模型", "混元模型格式、贴图或体积不适合 Godot Web。", "中", "中", "建立 GLB/GLTF 导入验收和体积阈值；先验 3 个模型。", "客户端/QA", "2026-06-29"],
    ["R005", "反馈", "反馈分散在群聊，无法转为研发任务。", "中", "高", "定义反馈字段、标签和周报模板；至少保留可复现上下文。", "产品/运营", "2026-06-28"],
    ["R006", "大赛", "官网投稿接口和账号方案晚于工具侧排期。", "高", "中", "先支持导出包 + manifest + 人工上传兜底；并行推进接口。", "产品/官网", "2026-07-10"],
    ["R007", "隐私安全", "上传日志、token 或项目路径可能带来隐私/安全问题。", "高", "中", "日志脱敏；token 安全存储；反馈上传前明确用户授权。", "客户端/安全", "2026-06-30"],
    ["R008", "发布", "更新清单、安装包、服务器地址不一致导致内测包无法升级。", "高", "中", "发布前检查 latest.yml/update.json/安装包 hash/固定下载链接。", "客户端/运维", "2026-06-29"],
    ["R009", "范围", "v0.2 同时承载资源、账号、大赛、反馈，范围过大。", "高", "高", "强制按 P0/P1/P2 切分；v0.2 只保核心闭环与手动导入。", "产品/项目管理", "2026-06-14"],
    ["R010", "多引擎", "过早抽象 Cocos/微信小游戏会拖慢 Godot 主线。", "中", "中", "6 月只做研究和接口草案，9 月再投入 POC。", "架构/产品", "2026-06-30"],
]

FIELD_ROWS = [
    ["字段", "说明"],
    ["优先级：高", "v0.2 核心闭环或内测/大赛关键阻塞，未完成会影响主要交付。"],
    ["优先级：中", "重要增强或跨线协同，需要排期但可根据核心链路让位。"],
    ["优先级：低", "优化、运营、反馈沉淀或不阻塞版本主目标的事项。"],
    ["状态：待开始", "尚未进入设计或开发。"],
    ["状态：进行中", "已经有实现或验证动作。"],
    ["状态：待验收", "基础实现已具备，需要内测/QA 结果确认。"],
    ["状态：已具备基础能力", "当前代码或流程已有能力，但仍可能需要真实场景打磨。"],
    ["状态：研究中/待方案/依赖方案", "需要先完成方案、接口或外部协同，再进入实现。"],
    ["完成定义", "需求完成必须有交付物、验收标准、可追踪证据；研发任务完成必须可被 QA 或产品复核。"],
]


def reset_generated_sheets(wb):
    for name in GENERATED_SHEETS:
        if name in wb.sheetnames:
            wb.remove(wb[name])


def title(ws, text, subtitle, max_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.cell(1, 1).value = text
    ws.cell(1, 1).font = Font(name="Microsoft YaHei", size=18, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="1F2937")
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    ws.cell(2, 1).value = subtitle
    ws.cell(2, 1).font = Font(name="Microsoft YaHei", size=10, color="374151")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor="E5E7EB")
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def apply_base(ws):
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color=cell.font.color.rgb if cell.font and cell.font.color and cell.font.color.type == "rgb" else "111827", bold=cell.font.bold if cell.font else False)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_table(ws, headers, rows, start_row, table_name, table_style="TableStyleMedium2"):
    thin = Side(style="thin", color="D1D5DB")
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="111827")
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row_idx, row in enumerate(rows, start_row + 1):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    end_row = start_row + len(rows)
    end_col = len(headers)
    ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name=table_style, showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)
    return end_row, end_col


def set_widths(ws, widths):
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_priority_status_formats(ws, priority_range, status_range=None):
    fills = {
        "高": PatternFill("solid", fgColor="FEE2E2"),
        "中": PatternFill("solid", fgColor="FEF3C7"),
        "低": PatternFill("solid", fgColor="E5E7EB"),
        "待验收": PatternFill("solid", fgColor="DBEAFE"),
        "进行中": PatternFill("solid", fgColor="D1FAE5"),
        "已具备基础能力": PatternFill("solid", fgColor="DCFCE7"),
        "规划中": PatternFill("solid", fgColor="F3E8FF"),
        "待方案": PatternFill("solid", fgColor="FFEDD5"),
        "研究中": PatternFill("solid", fgColor="EDE9FE"),
        "依赖方案": PatternFill("solid", fgColor="FCE7F3"),
        "依赖外部进度": PatternFill("solid", fgColor="FCE7F3"),
        "规划占位": PatternFill("solid", fgColor="E5E7EB"),
        "待开始": PatternFill("solid", fgColor="F3F4F6"),
    }
    for value in ["高", "中", "低"]:
        ws.conditional_formatting.add(priority_range, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fills[value]))
    if status_range:
        for value, fill in fills.items():
            if value in {"高", "中", "低"}:
                continue
            ws.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill))


def add_validations(ws, priority_range=None, status_range=None):
    if priority_range:
        dv = DataValidation(type="list", formula1='"高,中,低,暂缓"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(priority_range)
    if status_range:
        dv = DataValidation(type="list", formula1='"待开始,进行中,待验收,已具备基础能力,规划中,待方案,研究中,依赖方案,依赖外部进度,规划占位,待调研,暂缓"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(status_range)


def build_detail_sheet(wb):
    ws = wb.create_sheet("游戏创作平台roadmap细化版", 1)
    ws.sheet_properties.tabColor = "2563EB"
    title(
        ws,
        "万象游戏创作平台 Roadmap 细化版",
        "基于 2026-06-11 项目现状、0604 底稿和当前代码 README 梳理；原底稿保留在首个页签，本页用于版本评审与需求拆解。",
        10,
    )
    summary = [
        ["需求总数", f"=COUNTA(A9:A{8 + len(REQ_ROWS)})", "高优先级", f'=COUNTIF(E9:E{8 + len(REQ_ROWS)},"高")', "v0.2 需求", f'=COUNTIF(D9:D{8 + len(REQ_ROWS)},"*v0.2*")', "待验收/进行中", f'=COUNTIF(F9:F{8 + len(REQ_ROWS)},"待验收")+COUNTIF(F9:F{8 + len(REQ_ROWS)},"进行中")'],
        ["核心口径", "先保 Godot + AI CLI + Web 导出闭环", "资源口径", "v0.2 手动导入优先", "跨线口径", "账号/大赛先定协议", "后置口径", "Cocos/Chaos/手机样例 9 月规划"],
    ]
    for r, row in enumerate(summary, 4):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.fill = PatternFill("solid", fgColor="F9FAFB" if r == 4 else "EEF2FF")
            cell.border = Border(bottom=Side(style="thin", color="D1D5DB"))
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c % 2 == 1:
                cell.font = Font(name="Microsoft YaHei", bold=True, color="374151")
            else:
                cell.font = Font(name="Microsoft YaHei", color="111827")
        ws.row_dimensions[r].height = 28

    helper_start = 4
    ws.cell(helper_start, 12, "版本").font = Font(name="Microsoft YaHei", bold=True)
    ws.cell(helper_start, 13, "需求数").font = Font(name="Microsoft YaHei", bold=True)
    chart_rows = [
        ["v0.1 alpha", f'=COUNTIF($D$9:$D${8 + len(REQ_ROWS)},"*v0.1*")'],
        ["v0.2", f'=COUNTIF($D$9:$D${8 + len(REQ_ROWS)},"*v0.2*")'],
        ["研究", f'=COUNTIF($D$9:$D${8 + len(REQ_ROWS)},"*研究*")'],
        ["v0.4/规划", f'=COUNTIF($D$9:$D${8 + len(REQ_ROWS)},"*v0.4*")+COUNTIF($D$9:$D${8 + len(REQ_ROWS)},"*规划占位*")'],
    ]
    for idx, row in enumerate(chart_rows, helper_start + 1):
        ws.cell(idx, 12, row[0])
        ws.cell(idx, 13, row[1])
    chart = BarChart()
    chart.title = "需求按版本分布"
    chart.y_axis.title = "需求数"
    chart.x_axis.title = "版本"
    data = Reference(ws, min_col=13, min_row=helper_start, max_row=helper_start + len(chart_rows))
    cats = Reference(ws, min_col=12, min_row=helper_start + 1, max_row=helper_start + len(chart_rows))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 5.2
    chart.width = 8.5
    ws.add_chart(chart, "L10")

    headers = ["需求ID", "模块", "需求/能力", "版本/时间", "优先级", "当前状态", "范围拆解", "验收标准", "依赖/风险", "后续动作"]
    rows = [[r["id"], r["module"], r["name"], r["version"], r["priority"], r["status"], r["scope"], r["acceptance"], r["risk"], r["next"]] for r in REQ_ROWS]
    end_row, _ = write_table(ws, headers, rows, 8, "RoadmapDetailTable")
    set_widths(ws, [12, 18, 30, 26, 10, 14, 42, 42, 38, 38, 3, 16, 10])
    for row_idx in range(9, end_row + 1):
        ws.row_dimensions[row_idx].height = 76
    ws.freeze_panes = "A9"
    add_priority_status_formats(ws, f"E9:E{end_row}", f"F9:F{end_row}")
    add_validations(ws, f"E9:E{end_row}", f"F9:F{end_row}")
    return ws


def build_tasks_sheet(wb):
    ws = wb.create_sheet("后续需求任务池")
    ws.sheet_properties.tabColor = "16A34A"
    headers = ["任务ID", "关联需求ID", "版本", "阶段", "优先级", "任务名称", "任务说明", "交付物", "验收标准", "依赖/阻塞", "建议负责人/协作", "状态", "目标日期", "备注"]
    end_row, _ = write_table(ws, headers, TASK_ROWS, 1, "TaskPoolTable", "TableStyleMedium4")
    set_widths(ws, [10, 12, 14, 10, 10, 26, 44, 26, 42, 28, 18, 16, 14, 24])
    for row_idx in range(2, end_row + 1):
        ws.row_dimensions[row_idx].height = 58
    ws.freeze_panes = "A2"
    add_priority_status_formats(ws, f"E2:E{end_row}", f"L2:L{end_row}")
    add_validations(ws, f"E2:E{end_row}", f"L2:L{end_row}")
    return ws


def build_milestones_sheet(wb):
    ws = wb.create_sheet("版本里程碑")
    ws.sheet_properties.tabColor = "F59E0B"
    headers = ["里程碑ID", "目标日期", "版本/节点", "目标", "关键范围", "决策门", "退出标准"]
    end_row, _ = write_table(ws, headers, MILESTONES, 1, "MilestoneTable", "TableStyleMedium7")
    set_widths(ws, [12, 14, 24, 28, 48, 36, 42])
    for row_idx in range(2, end_row + 1):
        ws.row_dimensions[row_idx].height = 62
    ws.freeze_panes = "A2"
    return ws


def build_risks_sheet(wb):
    ws = wb.create_sheet("风险与依赖")
    ws.sheet_properties.tabColor = "DC2626"
    headers = ["风险ID", "领域", "风险/依赖", "影响", "概率", "缓解动作", "负责人/协作", "目标日期"]
    end_row, _ = write_table(ws, headers, RISKS, 1, "RiskTable", "TableStyleMedium3")
    set_widths(ws, [10, 14, 46, 10, 10, 50, 18, 14])
    for row_idx in range(2, end_row + 1):
        ws.row_dimensions[row_idx].height = 58
    ws.freeze_panes = "A2"
    add_priority_status_formats(ws, f"D2:D{end_row}", None)
    add_priority_status_formats(ws, f"E2:E{end_row}", None)
    return ws


def build_fields_sheet(wb):
    ws = wb.create_sheet("字段说明")
    ws.sheet_properties.tabColor = "6B7280"
    end_row, _ = write_table(ws, FIELD_ROWS[0], FIELD_ROWS[1:], 1, "FieldGuideTable", "TableStyleMedium1")
    set_widths(ws, [26, 90])
    for row_idx in range(2, end_row + 1):
        ws.row_dimensions[row_idx].height = 34
    ws.freeze_panes = "A2"
    return ws


def final_format(wb):
    for ws in wb.worksheets:
        if ws.title in GENERATED_SHEETS:
            ws.sheet_view.zoomScale = 85
        else:
            ws.sheet_view.showGridLines = False


def main():
    wb = load_workbook(SOURCE)
    reset_generated_sheets(wb)
    build_detail_sheet(wb)
    build_tasks_sheet(wb)
    build_milestones_sheet(wb)
    build_risks_sheet(wb)
    build_fields_sheet(wb)
    final_format(wb)
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
