from pathlib import Path
import sys

from openpyxl import load_workbook


sys.stdout.reconfigure(encoding="utf-8")
workbook_path = Path("outputs/requirements-roadmap/source.xlsx")
wb = load_workbook(workbook_path, data_only=False)

for ws in wb.worksheets:
    print(f"SHEET\t{ws.title}\trows={ws.max_row}\tcols={ws.max_column}")
    merged = [str(rng) for rng in ws.merged_cells.ranges]
    if merged:
        print("MERGED\t" + ", ".join(merged[:20]))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        values = ["" if value is None else str(value).replace("\n", " ") for value in row[: min(ws.max_column, 12)]]
        print("ROW\t" + "\t".join(values))
    print()
