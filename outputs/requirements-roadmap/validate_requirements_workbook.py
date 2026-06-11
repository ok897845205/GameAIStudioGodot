from pathlib import Path
import sys

from openpyxl import load_workbook


sys.stdout.reconfigure(encoding="utf-8")
path = Path("outputs/requirements-roadmap/0604_万象游戏创作平台roadmap_v1.2_需求拆解完成版.xlsx")
wb = load_workbook(path, data_only=False)

print(f"FILE\t{path.resolve()}")
print("SHEETS\t" + " | ".join(wb.sheetnames))

errors = []
formula_error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
for ws in wb.worksheets:
    non_empty = 0
    formula_errors = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                non_empty += 1
                if isinstance(cell.value, str) and any(token in cell.value for token in formula_error_tokens):
                    formula_errors.append(cell.coordinate)
    print(f"SHEET\t{ws.title}\trows={ws.max_row}\tcols={ws.max_column}\tnon_empty={non_empty}\ttables={len(ws.tables)}\tcharts={len(ws._charts)}")
    if ws.title != "游戏创作平台（底稿）" and non_empty < 2:
        errors.append(f"{ws.title} appears empty")
    if formula_errors:
        errors.append(f"{ws.title} formula error tokens at {formula_errors[:10]}")

detail = wb["游戏创作平台roadmap细化版"]
task = wb["后续需求任务池"]
req_ids = {detail.cell(row=r, column=1).value for r in range(9, detail.max_row + 1) if str(detail.cell(row=r, column=1).value or "").startswith("REQ-")}
task_req_ids = {task.cell(row=r, column=2).value for r in range(2, task.max_row + 1) if str(task.cell(row=r, column=1).value or "").startswith("T")}
missing = sorted(task_req_ids - req_ids)
print(f"DETAIL_REQS\t{len(req_ids)}")
print(f"TASK_ROWS\t{task.max_row - 1}")
print(f"TASK_REQ_MISSING\t{missing}")
if missing:
    errors.append(f"Task rows reference missing requirement IDs: {missing}")

if errors:
    print("VALIDATION\tFAIL")
    for error in errors:
        print(f"ERROR\t{error}")
    raise SystemExit(1)

print("VALIDATION\tOK")
